"""DOCX conversion module for FileFlow.
Handles:
- TXT -> DOCX (using python-docx)
"""

import io
from docx import Document
from docx.shared import Pt, Inches, RGBColor


def convert_text_to_docx(input_bytes: bytes, original_filename: str = "document") -> tuple[bytes, str, str]:
    """Convert raw text content to a Microsoft Word DOCX document.
    
    Args:
        input_bytes: Raw binary content of the TXT file
        original_filename: Name of the uploaded file
        
    Returns:
        tuple of (output_bytes, output_filename, mime_type)
    """
    stem = original_filename.rsplit(".", 1)[0] if "." in original_filename else original_filename

    try:
        try:
            text_str = input_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text_str = input_bytes.decode("latin-1")
    except Exception as e:
        raise ValueError(f"Unable to decode text content: {str(e)}")

    doc = Document()

    # Set standard 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Document Title heading
    title_p = doc.add_heading(original_filename, level=1)
    title_run = title_p.runs[0] if title_p.runs else title_p.add_run()
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)

    # Add text paragraphs
    lines = text_str.split("\n")
    current_para = []
    
    for line in lines:
        if line.strip() == "":
            if current_para:
                p = doc.add_paragraph(" ".join(current_para))
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(6)
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(11)
                current_para = []
            else:
                doc.add_paragraph("")
        else:
            current_para.append(line)

    if current_para:
        p = doc.add_paragraph(" ".join(current_para))
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(11)

    output_buffer = io.BytesIO()
    doc.save(output_buffer)
    output_filename = f"{stem}.docx"
    mime_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    return output_buffer.getvalue(), output_filename, mime_type
